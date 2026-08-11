#!/usr/bin/env python3
"""Fixed-Gear-racetag-Exporte ins SRB-„Amtliches Ergebnis"-Format überführen.

Erzeugt fixed-ergebnisse-srb.xlsx mit je einem Sheet pro Wertung im Layout
der Jedermann-Sheets aus "Karli Krit 08.08.26 Start und Ergebnislisten.xls":
Kopfblock (SRB / Amtliches Ergebnis / Veranstaltung / Ort, Datum / Runden),
Spalten Platz | St.-Nr. | Name, Vorname | Verein | Zeit | Runden.
Zeit = total_time_ms als echte Excel-Zeit (h:mm:ss); Verein/Team kommt
aus den Fixed-Rennen der racetag-DB (Portal-Teams, nachgetragen 11.08.).
"""
import csv
import io
import re
from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).parent
OUT = HERE / "fixed-ergebnisse-srb.xlsx"

# (CSV, [(Sheetname, Wertungstitel, bib_von, bib_bis)])
SOURCES = [
    ("racetag-19-30-fixed-gear-b-flinta-25-min-2026-08-08.csv",
     [("Fixed Gear B", "Fixed Gear B-Finale", 411, 436),
      ("FLINTA", "FLINTA*", 441, 446)]),
    ("racetag-20-00-fixed-gear-a-final-30-min-2026-08-08.csv",
     [("Fixed Gear A", "Fixed Gear A-Finale", 411, 436)]),
]


def read_export(path):
    raw = Path(path).read_text(encoding="utf-8-sig")
    comments = [l for l in raw.splitlines() if l.startswith("#")]
    body = "\n".join(l for l in raw.splitlines() if l and not l.startswith("#"))
    rows = list(csv.DictReader(io.StringIO(body)))
    total_laps = next((re.search(r"(\d+)", c).group(1) for c in comments
                       if "Total laps" in c), "")
    return rows, total_laps


def lade_vereine() -> dict:
    """bib -> Verein/Team aus den Fixed-Rennen (19:30 + 20:00) der racetag-DB."""
    import sqlite3
    db = Path.home() / ".racetag/data/racetag.db"
    con = sqlite3.connect(f"file:{db}?mode=ro", uri=True)
    out: dict = {}
    for pat in ("19:30%", "20:00%"):
        race = con.execute("SELECT id FROM races WHERE name LIKE ?", (pat,)).fetchone()
        if race:
            for bib, verein in con.execute(
                    "SELECT bib, verein FROM riders WHERE race_id=? AND verein != ''",
                    (race[0],)):
                out.setdefault(str(bib), verein)
    return out


VEREINE = lade_vereine()

wb = Workbook()
wb.remove(wb.active)
for src, sheets in SOURCES:
    rows, total_laps = read_export(HERE / src)
    for sheet_name, titel, lo, hi in sheets:
        ws = wb.create_sheet(sheet_name)
        ws["A1"] = "Sächsischer Radfahrer-Bund e. V."
        ws["A4"] = "Amtliches Ergebnis"
        ws["A6"] = "Karli Krit- 3. Lauf Revolution Crit"
        ws["F7"] = "Leipzig, 08.08.2026"
        ws["F8"] = "Ort, Datum"
        ws["E10"] = f"{total_laps} Runden"
        ws["A11"] = titel
        header = ["Platz", "St.-Nr.", "Name, Vorname", "Verein", "Zeit", "Runden"]
        for j, h in enumerate(header, 1):
            ws.cell(row=14, column=j, value=h)

        wertung = [r for r in rows
                   if r["bib"].isdigit() and lo <= int(r["bib"]) <= hi]
        out_row = 15
        platz = 0
        # Reihenfolge des Exports beibehalten (ist die Wertung), Status-Zeilen ans Ende
        for r in [x for x in wertung if not x["status"].strip()] + \
                 [x for x in wertung if x["status"].strip()]:
            status = r["status"].strip().upper()
            if status:
                ws.cell(row=out_row, column=1, value=status)
            else:
                platz += 1
                ws.cell(row=out_row, column=1, value=platz)
            ws.cell(row=out_row, column=2, value=int(r["bib"]))
            name = r["name"].strip()
            ws.cell(row=out_row, column=3, value=name)
            ws.cell(row=out_row, column=4, value=VEREINE.get(r["bib"].strip(), ""))
            if r["total_time_ms"].strip():
                c = ws.cell(row=out_row, column=5,
                            value=int(r["total_time_ms"]) / 86400000.0)
                c.number_format = "h:mm:ss"
            ws.cell(row=out_row, column=6, value=int(r["laps"] or 0))
            out_row += 1
        print(f"{sheet_name}: {out_row - 15} Zeilen ({titel}, {total_laps} Runden)")

wb.save(OUT)
print(f"-> {OUT}")
