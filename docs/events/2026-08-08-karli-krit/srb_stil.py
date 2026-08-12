"""SRB-Vorlagen-Formatierung — 1:1 aus Kerstins xls ausgelesen (12.08.).

Vermessen an 'Karli Krit 08.08.26 Start und Ergebnislisten.xls', Blatt
'2.1 Masters 4' (xlrd, formatting_info=True):
  A1  Bookman Old Style 26 fett, zentriert, verbunden A1:F1
  A4  Arial Narrow 20 fett unterstrichen, zentriert, verbunden A4:F4
  A6  Arial 13 fett, zentriert, verbunden A6:F6
  F7/F8  Arial 10 (Ort, Datum)
  D10 'N Runden = N km', verbunden D10:E11, linksbündig
  A11 Klasse, Arial 10, zentriert, verbunden A11:C11
  Z.14 Tabellenkopf Arial Narrow 12 fett, dünner Rahmen ringsum
  Datenzeilen: Platz-Spalte Arial 11 fett mit Rahmen, Rest Arial 10
  Spaltenbreiten (Zeichen): 5.5 / 6 / 22.5 / 32 / 14 / 9 / 9
"""
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

_DUENN = Side(style="thin")
RAHMEN = Border(top=_DUENN, bottom=_DUENN, left=_DUENN, right=_DUENN)
BREITEN = [5.5, 6, 22.5, 32, 14, 9, 9]
ZENTRIERT = Alignment(horizontal="center")


def formatiere(ws, spalten: int) -> None:
    """Vorlagen-Formatierung auf ein fertig befülltes Blatt anwenden.

    Erwartet das Standard-Layout: Kopfblock A1/A4/A6/F7/F8/D10/A11,
    Tabellenkopf in Zeile 14, Daten ab Zeile 15 (Hinweise darunter).
    """
    ende = get_column_letter(spalten)
    for i, breite in enumerate(BREITEN[:max(spalten, 7)], 1):
        ws.column_dimensions[get_column_letter(i)].width = breite

    for zelle, ab, bis in (("A1", "A1", f"{ende}1"), ("A4", "A4", f"{ende}4"),
                           ("A6", "A6", f"{ende}6")):
        ws.merge_cells(f"{ab}:{bis}")
        ws[zelle].alignment = ZENTRIERT
    ws["A1"].font = Font(name="Bookman Old Style", size=26, bold=True)
    ws["A4"].font = Font(name="Arial Narrow", size=20, bold=True,
                         underline="single")
    ws["A6"].font = Font(name="Arial", size=13, bold=True)
    ws["F7"].font = Font(name="Arial", size=10)
    ws["F8"].font = Font(name="Arial", size=10)
    ws.merge_cells("D10:E11")
    ws["D10"].alignment = Alignment(horizontal="left", vertical="center")
    ws["D10"].font = Font(name="Arial", size=10)
    ws.merge_cells("A11:C11")
    ws["A11"].alignment = ZENTRIERT
    ws["A11"].font = Font(name="Arial", size=10)

    kopf = Font(name="Arial Narrow", size=12, bold=True)
    for j in range(1, spalten + 1):
        c = ws.cell(row=14, column=j)
        c.font = kopf
        c.border = RAHMEN
        if j == 1:
            c.alignment = ZENTRIERT

    for zeile in range(15, ws.max_row + 1):
        platz = ws.cell(row=zeile, column=1)
        wert = platz.value
        if wert is None:
            continue
        if isinstance(wert, str) and wert.startswith("Hinweis:"):
            platz.font = Font(name="Arial", size=9, italic=True)
            continue
        platz.font = Font(name="Arial", size=11, bold=True)
        platz.border = RAHMEN
        platz.alignment = ZENTRIERT
        for j in range(2, spalten + 1):
            ws.cell(row=zeile, column=j).font = Font(name="Arial", size=10)
