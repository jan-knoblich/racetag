#!/usr/bin/env python3
"""Lauf-Ergebnisse als editierbare Excel-Arbeitsmappe (Vorlagen-Layout).

Sechs Wertungen aus den ergebnis-lauf-*.csv als je ein Blatt, Kopfblock wie
die SRB-Vorlage (Distanz im D10-Feld), Vereine aus der racetag-DB, Zeiten
als echte Excel-Zeiten. Nicht Gewertete stehen mit „–" und Hinweis unten.

Usage: e2e-venv/bin/python lauf_xlsx.py  ->  ~/Downloads/karli-lauf-ergebnisse.xlsx
"""
import csv
import sqlite3
from pathlib import Path

from openpyxl import Workbook

from srb_stil import formatiere

HERE = Path(__file__).parent
OUT = Path.home() / "Downloads/karli-lauf-ergebnisse.xlsx"

LAUF = [  # (Datei, Blattname, Klasse, Distanz)
    ("ergebnis-lauf-10km-erwachsene.csv", "10km", "Lauf 10 km Erwachsene", "13 Runden = 10 km"),
    ("ergebnis-lauf-10km-u18-maennlich.csv", "10km U18m", "Lauf 10 km U18 männlich", "13 Runden = 10 km"),
    ("ergebnis-lauf-10km-u18-weiblich.csv", "10km U18w", "Lauf 10 km U18 weiblich", "13 Runden = 10 km"),
    ("ergebnis-lauf-5km-erwachsene.csv", "5km", "Lauf 5 km Erwachsene", "7 Runden = 5 km"),
    ("ergebnis-lauf-5km-u18-maennlich.csv", "5km U18m", "Lauf 5 km U18 männlich", "7 Runden = 5 km"),
    ("ergebnis-lauf-5km-u18-weiblich.csv", "5km U18w", "Lauf 5 km U18 weiblich", "7 Runden = 5 km"),
]


def lauf_vereine() -> dict[str, str]:
    con = sqlite3.connect(f"file:{Path.home()}/.racetag/data/racetag.db?mode=ro", uri=True)
    race = con.execute("SELECT id FROM races WHERE name LIKE '%Lauf%'").fetchone()
    return {str(b): v for b, v in con.execute(
        "SELECT bib, verein FROM riders WHERE race_id=? AND verein != ''", (race[0],))}


def zeit_als_excel(hms: str):
    teile = hms.split(":")
    if len(teile) != 3:
        return hms
    h, m, s = (int(t) for t in teile)
    return (h * 3600 + m * 60 + s) / 86400.0


wb = Workbook()
wb.remove(wb.active)
vereine = lauf_vereine()

for datei, blatt, klasse, distanz in LAUF:
    ws = wb.create_sheet(blatt)
    ws["A1"] = "Karli Lauf — Revolution Crit"
    ws["A4"] = "Amtliches Ergebnis"
    ws["A6"] = "Karli Krit + Karli Lauf"
    ws["F7"] = "Leipzig, 08.08.2026"
    ws["F8"] = "Ort, Datum"
    ws["D10"] = distanz
    ws["A11"] = klasse
    for j, h in enumerate(["Platz", "St.-Nr.", "Name, Vorname", "Verein",
                           "Zeit", "Runden", "Hinweis"], 1):
        ws.cell(row=14, column=j, value=h)
    out_row = 15
    with open(HERE / datei, encoding="utf-8-sig") as f:
        for r in csv.DictReader(f, delimiter=";"):
            ws.cell(row=out_row, column=1,
                    value=int(r["platz"]) if r["platz"].isdigit() else "–")
            ws.cell(row=out_row, column=2, value=int(r["nummer"]))
            ws.cell(row=out_row, column=3, value=r["name"])
            ws.cell(row=out_row, column=4, value=vereine.get(r["nummer"], ""))
            if r["zeit"]:
                c = ws.cell(row=out_row, column=5, value=zeit_als_excel(r["zeit"]))
                c.number_format = "h:mm:ss"
            ws.cell(row=out_row, column=6, value=int(r["runden"]))
            ws.cell(row=out_row, column=7, value=r["hinweis"])
            out_row += 1
    formatiere(ws, 7)
    ws.column_dimensions["G"].width = 44  # Hinweis-Spalte lesbar

wb.save(OUT)
print(f"{OUT}  ({len(wb.sheetnames)} Blätter)")
