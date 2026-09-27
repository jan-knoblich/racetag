"""Excel workbook builders (docs/PLAN-EXCEL-EXPORT.md).

Two exports, both German, both meant to be opened by double-click:

* ``build_results_workbook`` — the result of one race for the operator.
* ``build_readings_workbook`` — every stored reading for the data check.

The builders are pure: they take plain dicts and return ``bytes``. They never
import ``app`` (the desktop shell loads the backend under a different module
name) and never touch the database; the caller assembles the rows.
"""
from __future__ import annotations

import io
import re
import unicodedata
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

XLSX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

# German Excel display formats. Excel itself renders them with the user's
# locale separators, so a German Windows shows "16.09.2026 14:03:12,500".
FMT_DATETIME = "DD.MM.YYYY HH:MM:SS.000"
FMT_DURATION = "[h]:mm:ss.0"
FMT_SECONDS = "0.000"

_HEADER_FONT = Font(bold=True)
_HEADER_ALIGN = Alignment(vertical="center")

# Excel's epoch (the 1900 date system openpyxl writes).
_EXCEL_EPOCH = datetime(1899, 12, 30)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def sanitize_filename(stem: str, suffix: str = ".xlsx") -> str:
    """ASCII filename for Content-Disposition: umlauts transliterated, spaces
    to dashes. Windows Explorer and every mail client cope with that."""
    replacements = {"ä": "ae", "ö": "oe", "ü": "ue", "Ä": "Ae", "Ö": "Oe",
                    "Ü": "Ue", "ß": "ss"}
    for src, dst in replacements.items():
        stem = stem.replace(src, dst)
    stem = unicodedata.normalize("NFKD", stem).encode("ascii", "ignore").decode()
    stem = re.sub(r"[^A-Za-z0-9._-]+", "-", stem)
    stem = re.sub(r"-{2,}", "-", stem).strip("-._") or "export"
    return f"{stem[:120]}{suffix}"


def parse_iso(value: Optional[str]) -> Optional[datetime]:
    """Parse the backend's ISO-8601 UTC strings; return an aware datetime."""
    if not value:
        return None
    text = value.strip().replace("Z", "+00:00")
    try:
        parsed = datetime.fromisoformat(text)
    except ValueError:
        return None
    if parsed.tzinfo is None:
        return parsed.replace(tzinfo=timezone.utc)
    return parsed.astimezone(timezone.utc)


def utc_cell(value: Optional[str]) -> Optional[datetime]:
    """Naive UTC datetime for Excel (Excel has no timezone concept)."""
    parsed = parse_iso(value)
    return None if parsed is None else parsed.replace(tzinfo=None)


def local_cell(value: Optional[str], tz=None) -> Optional[datetime]:
    """Same instant in the machine's local timezone, naive, for Excel."""
    parsed = parse_iso(value)
    if parsed is None:
        return None
    return parsed.astimezone(tz).replace(tzinfo=None)


def duration_cell(seconds: Optional[float]) -> Optional[float]:
    """Excel duration (a fraction of a day) for the [h]:mm:ss.0 format."""
    if seconds is None:
        return None
    return float(seconds) / 86400.0


def ms_duration_cell(millis: Optional[int]) -> Optional[float]:
    if millis is None:
        return None
    return float(millis) / 1000.0 / 86400.0


def _write_sheet(
    ws: Worksheet,
    headers: Sequence[str],
    rows: List[Sequence[Any]],
    *,
    widths: Optional[Sequence[int]] = None,
    formats: Optional[Dict[int, str]] = None,
    autofilter: bool = True,
) -> None:
    """Write a header row plus data, then apply the usual comforts."""
    ws.append(list(headers))
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.alignment = _HEADER_ALIGN
    for row in rows:
        ws.append(list(row))

    if formats:
        for col_index, number_format in formats.items():
            letter = get_column_letter(col_index)
            for (cell,) in ws.iter_rows(
                min_row=2, min_col=col_index, max_col=col_index
            ):
                cell.number_format = number_format

    for index, header in enumerate(headers, start=1):
        width = widths[index - 1] if widths and index - 1 < len(widths) else None
        ws.column_dimensions[get_column_letter(index)].width = width or max(
            12, min(40, len(str(header)) + 4)
        )

    ws.freeze_panes = "A2"
    if autofilter and rows:
        ws.auto_filter.ref = (
            f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
        )


def _meta_sheet(ws: Worksheet, pairs: List[Sequence[Any]]) -> None:
    ws.append(["Feld", "Wert"])
    for cell in ws[1]:
        cell.font = _HEADER_FONT
    for key, value in pairs:
        ws.append([key, value])
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 44
    for (cell,) in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        if isinstance(cell.value, datetime):
            cell.number_format = FMT_DATETIME


def _workbook_bytes(wb: Workbook) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _status_text(status: Optional[str]) -> str:
    return (status or "").upper()


def _yes_no(value: bool) -> str:
    return "ja" if value else "nein"


# ---------------------------------------------------------------------------
# Results workbook
# ---------------------------------------------------------------------------

RESULT_HEADERS = [
    "Platz", "Startnummer", "Name", "Verein", "UCI-ID", "Runden",
    "Gesamtzeit", "Netto-Zeit", "Letzte Durchfahrt", "Rückstand (Runden)",
    "Status", "Tag-ID",
]
_RESULT_WIDTHS = [7, 12, 26, 24, 16, 8, 14, 14, 22, 18, 9, 34]

RIDER_HEADERS = [
    "Startnummer", "Name", "Verein", "UCI-ID", "Tag-ID", "Status", "Angelegt",
]
_RIDER_WIDTHS = [12, 26, 24, 16, 34, 9, 22]


def _rider_rows(riders: List[Dict[str, Any]], tz=None) -> List[List[Any]]:
    def sort_key(rider: Dict[str, Any]):
        bib = str(rider.get("bib") or "")
        return (0, int(bib), "") if bib.isdigit() else (1, 0, bib or rider.get("tag_id") or "")

    rows: List[List[Any]] = []
    for rider in sorted(riders, key=sort_key):
        rows.append([
            rider.get("bib") or "",
            rider.get("name") or "",
            rider.get("verein") or "",
            rider.get("uci_id") or "",
            rider.get("tag_id") or "",
            _status_text(rider.get("status")),
            local_cell(rider.get("created_at"), tz),
        ])
    return rows


def build_results_workbook(
    *,
    race: Dict[str, Any],
    standings: List[Dict[str, Any]],
    riders: List[Dict[str, Any]],
    exported_at: Optional[datetime] = None,
    version: Optional[str] = None,
    tz=None,
) -> bytes:
    """One race: Ergebnis, Rennen, Fahrer."""
    exported_at = exported_at or datetime.now(timezone.utc)
    wb = Workbook()

    ws = wb.active
    ws.title = "Ergebnis"
    rows: List[List[Any]] = []
    position = 0
    for item in standings:
        status = _status_text(item.get("status"))
        if not status:
            position += 1
        rows.append([
            "" if status else position,
            item.get("bib") or "",
            item.get("name") or "",
            item.get("verein") or "",
            item.get("uci_id") or "",
            item.get("laps", 0),
            ms_duration_cell(item.get("total_time_ms")),
            ms_duration_cell(item.get("net_time_ms")),
            local_cell(item.get("last_pass_time"), tz),
            item.get("laps_behind") if item.get("laps_behind") is not None else "",
            status,
            item.get("tag_id") or "",
        ])
    _write_sheet(
        ws, RESULT_HEADERS, rows,
        widths=_RESULT_WIDTHS,
        formats={7: FMT_DURATION, 8: FMT_DURATION, 9: FMT_DATETIME},
    )

    _meta_sheet(wb.create_sheet("Rennen"), [
        ["Rennen", race.get("name") or ""],
        ["Geplant", local_cell(race.get("scheduled_at"), tz)],
        ["Gestartet", local_cell(race.get("started_at"), tz)],
        ["Beendet", local_cell(race.get("ended_at"), tz)],
        ["Runden (Ziel)", race.get("total_laps")],
        ["Wertungsmodus", race.get("finish_mode") or ""],
        ["Dauer (s)", race.get("duration_s")],
        ["Mindest-Rundenzeit (s)", race.get("min_pass_interval_s")],
        ["Fahrer in der Wertung", len(standings)],
        ["Gekoppelte Fahrer", len(riders)],
        ["Racetag-Version", version or ""],
        ["Exportiert am", exported_at.astimezone(tz).replace(tzinfo=None)],
        ["Zeitangaben", "Ortszeit dieses Rechners"],
    ])

    _write_sheet(
        wb.create_sheet("Fahrer"), RIDER_HEADERS, _rider_rows(riders, tz),
        widths=_RIDER_WIDTHS, formats={7: FMT_DATETIME},
    )
    return _workbook_bytes(wb)


# ---------------------------------------------------------------------------
# Readings workbook
# ---------------------------------------------------------------------------

READING_HEADERS = [
    "Nr", "Rennen", "Zeit (lokal)", "Zeit (UTC)", "Sekunden seit Start",
    "Tag-ID", "Startnummer", "Name", "Verein", "Antenne", "RSSI", "Ereignis",
    "Reader-Seriennummer", "Gewertet", "Runde",
    "Abstand zur vorherigen Lesung (s)",
    "Abstand zur vorherigen gewerteten Runde (s)",
]
_READING_WIDTHS = [7, 24, 22, 22, 18, 34, 12, 24, 22, 9, 9, 10, 20, 10, 8, 18, 18]

SUMMARY_HEADERS = [
    "Rennen", "Startnummer", "Name", "Tag-ID", "Lesungen", "Gewertet",
    "Verworfen", "Erste Lesung", "Letzte Lesung", "Antennen",
]
_SUMMARY_WIDTHS = [24, 12, 24, 34, 10, 10, 10, 22, 22, 16]

RACE_HEADERS = [
    "Rennen", "Geplant", "Gestartet", "Beendet", "Runden (Ziel)",
    "Mindest-Rundenzeit (s)", "Lesungen", "Gekoppelte Fahrer", "Rennen-ID",
]
_RACE_WIDTHS = [24, 22, 22, 22, 14, 20, 12, 18, 38]

ALL_RIDER_HEADERS = ["Rennen"] + RIDER_HEADERS


def build_readings_workbook(
    *,
    races: List[Dict[str, Any]],
    readings: List[Dict[str, Any]],
    riders_by_race: Dict[str, List[Dict[str, Any]]],
    exported_at: Optional[datetime] = None,
    version: Optional[str] = None,
    tz=None,
) -> bytes:
    """Every stored reading: Lesungen, Zusammenfassung, Rennen, Fahrer, Info.

    ``readings`` must already be ordered (race, then timestamp) and carry the
    reconstruction fields (``counted``, ``lap``, ``gap_previous_s``,
    ``gap_previous_lap_s``, ``seconds_since_start``).
    """
    exported_at = exported_at or datetime.now(timezone.utc)
    wb = Workbook()

    ws = wb.active
    ws.title = "Lesungen"
    rows: List[List[Any]] = []
    for index, reading in enumerate(readings, start=1):
        rows.append([
            index,
            reading.get("race_name") or "",
            local_cell(reading.get("timestamp"), tz),
            utc_cell(reading.get("timestamp")),
            reading.get("seconds_since_start"),
            reading.get("tag_id") or "",
            reading.get("bib") or "",
            reading.get("name") or "",
            reading.get("verein") or "",
            reading.get("antenna"),
            reading.get("rssi"),
            reading.get("event_type") or "",
            reading.get("reader_serial") or "",
            _yes_no(bool(reading.get("counted"))),
            reading.get("lap") if reading.get("lap") else "",
            reading.get("gap_previous_s"),
            reading.get("gap_previous_lap_s"),
        ])
    _write_sheet(
        ws, READING_HEADERS, rows,
        widths=_READING_WIDTHS,
        formats={
            3: FMT_DATETIME, 4: FMT_DATETIME, 5: FMT_SECONDS,
            16: FMT_SECONDS, 17: FMT_SECONDS,
        },
    )
    ws["N1"].comment = _counted_comment()

    # ---- Zusammenfassung ---------------------------------------------------
    summary: Dict[tuple, Dict[str, Any]] = {}
    for reading in readings:
        key = (reading.get("race_id"), reading.get("tag_id"))
        entry = summary.setdefault(key, {
            "race_name": reading.get("race_name") or "",
            "bib": reading.get("bib") or "",
            "name": reading.get("name") or "Unbekannt",
            "tag_id": reading.get("tag_id") or "",
            "reads": 0, "counted": 0, "first": None, "last": None,
            "antennas": {},
        })
        entry["reads"] += 1
        if reading.get("counted"):
            entry["counted"] += 1
        stamp = reading.get("timestamp")
        if stamp and (entry["first"] is None or stamp < entry["first"]):
            entry["first"] = stamp
        if stamp and (entry["last"] is None or stamp > entry["last"]):
            entry["last"] = stamp
        antenna = reading.get("antenna")
        label = str(antenna) if antenna is not None else "?"
        entry["antennas"][label] = entry["antennas"].get(label, 0) + 1

    summary_rows = []
    for entry in sorted(
        summary.values(),
        key=lambda e: (e["race_name"], -e["reads"], e["tag_id"]),
    ):
        antennas = ", ".join(
            f"{ant}: {count}" for ant, count in sorted(entry["antennas"].items())
        )
        summary_rows.append([
            entry["race_name"], entry["bib"], entry["name"], entry["tag_id"],
            entry["reads"], entry["counted"], entry["reads"] - entry["counted"],
            local_cell(entry["first"], tz), local_cell(entry["last"], tz),
            antennas,
        ])
    _write_sheet(
        wb.create_sheet("Zusammenfassung"), SUMMARY_HEADERS, summary_rows,
        widths=_SUMMARY_WIDTHS, formats={8: FMT_DATETIME, 9: FMT_DATETIME},
    )

    # ---- Rennen ------------------------------------------------------------
    race_rows = []
    for race in races:
        race_rows.append([
            race.get("name") or "",
            local_cell(race.get("scheduled_at"), tz),
            local_cell(race.get("started_at"), tz),
            local_cell(race.get("ended_at"), tz),
            race.get("total_laps"),
            race.get("min_pass_interval_s"),
            race.get("reading_count"),
            len(riders_by_race.get(race.get("id"), [])),
            race.get("id") or "",
        ])
    _write_sheet(
        wb.create_sheet("Rennen"), RACE_HEADERS, race_rows,
        widths=_RACE_WIDTHS,
        formats={2: FMT_DATETIME, 3: FMT_DATETIME, 4: FMT_DATETIME},
    )

    # ---- Fahrer ------------------------------------------------------------
    rider_rows: List[List[Any]] = []
    for race in races:
        name = race.get("name") or ""
        for row in _rider_rows(riders_by_race.get(race.get("id"), []), tz):
            rider_rows.append([name] + row)
    _write_sheet(
        wb.create_sheet("Fahrer"), ALL_RIDER_HEADERS, rider_rows,
        widths=[24] + _RIDER_WIDTHS, formats={8: FMT_DATETIME},
    )

    # ---- Info --------------------------------------------------------------
    _meta_sheet(wb.create_sheet("Info"), [
        ["Inhalt", "Jede gespeicherte Lesung, auch verworfene"],
        ["Rennen im Export", len(races)],
        ["Lesungen gesamt", len(readings)],
        ["Davon gewertet", sum(1 for r in readings if r.get("counted"))],
        ["Gewertet / Runde", "rekonstruiert (siehe Kommentar in „Lesungen“)"],
        ["Zeitangaben", "„Zeit (lokal)“ = Ortszeit dieses Rechners, „Zeit (UTC)“ = Reader-Zeit"],
        ["Racetag-Version", version or ""],
        ["Exportiert am", exported_at.astimezone(tz).replace(tzinfo=None)],
    ])
    return _workbook_bytes(wb)


def _counted_comment():
    from openpyxl.comments import Comment

    return Comment(
        "Rekonstruiert: Racetag speichert jede Lesung, aber nicht, ob sie "
        "damals als Runde gezählt wurde. Diese Spalte spielt die Lesungen mit "
        "der Mindest-Rundenzeit und dem Startzeitpunkt des Rennens noch einmal "
        "durch. Manuell nachgetragene oder entfernte Runden sind darin nicht "
        "enthalten.",
        "Racetag",
        height=150,
        width=360,
    )
