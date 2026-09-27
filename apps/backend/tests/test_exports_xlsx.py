"""Excel exports (docs/PLAN-EXCEL-EXPORT.md).

Two workbooks: the race result for the operator and every stored reading for
the data check. The readings sheet must never hide a row — that file is what
Jan uses to repair couplings and missed passes after a race.
"""
import importlib
import io
from datetime import datetime, timedelta, timezone

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
START = datetime(2026, 9, 24, 10, 0, 0, tzinfo=timezone.utc)


def _iso(dt: datetime) -> str:
    return dt.isoformat(timespec="milliseconds").replace("+00:00", "Z")


@pytest.fixture()
def app_client(tmp_path, monkeypatch):
    monkeypatch.setenv("RACETAG_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("RACETAG_VERSION", "9.9.9")
    monkeypatch.delenv("RACETAG_API_KEY", raising=False)
    import app as app_module
    importlib.reload(app_module)
    with TestClient(app_module.app) as client:
        yield client, app_module


def _sheet_rows(payload: bytes, sheet: str):
    wb = load_workbook(io.BytesIO(payload))
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    return wb, ws, rows[0], rows[1:]


def _post_arrive(client, tag, when, antenna=1, rssi=-52):
    body = {"events": [{
        "source": "test", "reader_ip": "192.168.178.22", "reader_serial": "00179E01",
        "timestamp": _iso(when), "event_type": "arrive", "tag_id": tag,
        "antenna": antenna, "rssi": rssi,
    }]}
    res = client.post("/events/tag/batch", json=body)
    assert res.status_code == 200, res.text
    return res


def _seed(client, app_module):
    """A race with two registered riders, one unknown tag, a cooldown
    duplicate, two antennas, and one rider set to DNF."""
    client.patch("/config", json={"min_lap_interval_s": 8.0, "total_laps": 3})
    client.post("/riders", json={"tag_id": "AAA1", "bib": "7", "name": "Jürgen Öhler"})
    client.post("/riders", json={"tag_id": "BBB2", "bib": "12", "name": "Zoe Zander"})
    client.post("/race/start")
    # Anchor the start in memory AND in storage so the live gate and the
    # export reconstruction agree; otherwise the fixed event timestamps below
    # would predate the wall-clock start.
    race_id = client.get("/race").json()["id"]
    app_module.race.started_at = START
    app_module.storage.update_race(race_id, started_at=START)

    # Lap 1 for both, well after the start so the first-pass cooldown is met.
    _post_arrive(client, "AAA1", START + timedelta(seconds=30), antenna=1)
    _post_arrive(client, "BBB2", START + timedelta(seconds=32), antenna=2)
    # Duplicate inside the cooldown: stored, must NOT count as a lap.
    _post_arrive(client, "AAA1", START + timedelta(seconds=33), antenna=2)
    # Lap 2 for the first rider.
    _post_arrive(client, "AAA1", START + timedelta(seconds=70), antenna=1)
    # A tag nobody coupled: stored, never a lap, must appear in the export.
    _post_arrive(client, "CCC3", START + timedelta(seconds=75), antenna=2)

    client.patch("/riders/BBB2/status", json={"status": "dnf"})
    return app_module


def test_results_workbook_matches_the_standings(app_client):
    client, app_module = app_client
    _seed(client, app_module)

    res = client.get("/results.xlsx")
    assert res.status_code == 200
    assert res.headers["content-type"].startswith(XLSX_MIME)
    assert "Racetag-Ergebnis" in res.headers["content-disposition"]
    # Umlauts are transliterated for Content-Disposition.
    assert "ü" not in res.headers["content-disposition"]

    wb, ws, header, rows = _sheet_rows(res.content, "Ergebnis")
    assert wb.sheetnames == ["Ergebnis", "Rennen", "Fahrer"]
    assert header[:6] == ("Platz", "Startnummer", "Name", "Verein", "UCI-ID", "Runden")
    assert ws.freeze_panes == "A2"

    standings = client.get("/classification").json()["standings"]
    assert len(rows) == len(standings)
    by_bib = {r[1]: r for r in rows}
    assert by_bib["7"][2] == "Jürgen Öhler"
    assert by_bib["7"][5] == 2  # two counted laps
    # A DNF rider carries the status and no finishing position.
    assert by_bib["12"][10] == "DNF"
    assert by_bib["12"][0] is None  # openpyxl reads a blank cell as None
    # The classified rider does get one.
    assert by_bib["7"][0] == 1


def test_results_workbook_has_every_rider_even_without_readings(app_client):
    client, app_module = app_client
    _seed(client, app_module)
    client.post("/riders", json={"tag_id": "DDD4", "bib": "99", "name": "Ohne Lesung"})

    _, _, header, rows = _sheet_rows(client.get("/results.xlsx").content, "Fahrer")
    assert header[0] == "Startnummer"
    assert "99" in {r[0] for r in rows}


def test_readings_workbook_keeps_every_stored_reading(app_client):
    client, app_module = app_client
    _seed(client, app_module)

    res = client.get("/readings.xlsx")
    assert res.status_code == 200
    wb, ws, header, rows = _sheet_rows(res.content, "Lesungen")
    assert wb.sheetnames == ["Lesungen", "Zusammenfassung", "Rennen", "Fahrer", "Info"]

    # Five readings were posted, including the duplicate and the unknown tag.
    assert len(rows) == 5
    tags = [r[5] for r in rows]
    assert tags.count("AAA1") == 3
    assert "CCC3" in tags

    cols = {name: i for i, name in enumerate(header)}
    counted = [r[cols["Gewertet"]] for r in rows]
    assert counted.count("ja") == 3  # two for AAA1, one for BBB2
    # The cooldown duplicate is stored but not counted.
    dupes = [r for r in rows if r[cols["Tag-ID"]] == "AAA1"]
    assert [d[cols["Gewertet"]] for d in dupes] == ["ja", "nein", "ja"]
    assert [d[cols["Runde"]] for d in dupes] == [1, None, 2]
    # The unknown tag is present, without bib or name.
    unknown = next(r for r in rows if r[cols["Tag-ID"]] == "CCC3")
    assert unknown[cols["Startnummer"]] is None
    assert unknown[cols["Gewertet"]] == "nein"

    # Times are real datetimes, not text.
    assert isinstance(rows[0][cols["Zeit (UTC)"]], datetime)
    assert rows[0][cols["Zeit (UTC)"]] == START.replace(tzinfo=None) + timedelta(seconds=30)
    assert rows[0][cols["Sekunden seit Start"]] == pytest.approx(30.0)
    assert ws.cell(row=1, column=cols["Gewertet"] + 1).comment is not None

    # Antenna and RSSI stay numeric so Excel can filter on them.
    assert rows[0][cols["Antenne"]] == 1
    assert rows[0][cols["RSSI"]] == pytest.approx(-52)


def test_readings_counted_column_matches_the_api_lap_count(app_client):
    client, app_module = app_client
    _seed(client, app_module)

    _, _, header, rows = _sheet_rows(client.get("/readings.xlsx").content, "Lesungen")
    cols = {name: i for i, name in enumerate(header)}
    per_tag = {}
    for row in rows:
        if row[cols["Gewertet"]] == "ja":
            per_tag[row[cols["Tag-ID"]]] = per_tag.get(row[cols["Tag-ID"]], 0) + 1

    for item in client.get("/classification").json()["standings"]:
        assert per_tag.get(item["tag_id"], 0) == item["laps"]


def test_readings_summary_counts_discarded_readings(app_client):
    client, app_module = app_client
    _seed(client, app_module)

    _, _, header, rows = _sheet_rows(client.get("/readings.xlsx").content, "Zusammenfassung")
    cols = {name: i for i, name in enumerate(header)}
    by_tag = {r[cols["Tag-ID"]]: r for r in rows}
    assert by_tag["AAA1"][cols["Lesungen"]] == 3
    assert by_tag["AAA1"][cols["Gewertet"]] == 2
    assert by_tag["AAA1"][cols["Verworfen"]] == 1
    assert by_tag["CCC3"][cols["Name"]] == "Unbekannt"
    assert "1: 2" in by_tag["AAA1"][cols["Antennen"]]


def test_readings_per_race_and_all_races(app_client):
    client, app_module = app_client
    _seed(client, app_module)
    race_id = client.get("/race").json()["id"]

    single = client.get(f"/races/{race_id}/readings.xlsx")
    assert single.status_code == 200
    _, _, _, rows_single = _sheet_rows(single.content, "Lesungen")

    # A second race with its own reading; "all races" must contain both.
    second = client.post("/races", json={"name": "Lauf 2", "total_laps": 2}).json()
    client.post(f"/races/{second['id']}/activate")
    client.post("/riders", json={"tag_id": "EEE5", "bib": "3", "name": "Zweites Rennen"})
    client.post("/race/start")
    app_module.race.started_at = START + timedelta(hours=1)
    app_module.storage.update_race(second["id"], started_at=START + timedelta(hours=1))
    _post_arrive(client, "EEE5", START + timedelta(hours=1, seconds=30))

    _, _, header, rows_all = _sheet_rows(client.get("/readings.xlsx").content, "Lesungen")
    cols = {name: i for i, name in enumerate(header)}
    assert len(rows_all) == len(rows_single) + 1
    assert {r[cols["Rennen"]] for r in rows_all} == {"Default race", "Lauf 2"}

    # The per-race export works for a race that is no longer active.
    again = client.get(f"/races/{race_id}/readings.xlsx")
    assert again.status_code == 200
    _, _, _, rows_again = _sheet_rows(again.content, "Lesungen")
    assert len(rows_again) == len(rows_single)


def test_readings_workbook_for_an_empty_race(app_client):
    client, app_module = app_client
    res = client.get("/readings.xlsx")
    assert res.status_code == 200
    wb, _, _, rows = _sheet_rows(res.content, "Lesungen")
    assert rows == []
    assert wb["Info"]["B4"].value == 0  # Lesungen gesamt


def test_unknown_race_is_404_and_inactive_results_are_409(app_client):
    client, app_module = app_client
    _seed(client, app_module)
    assert client.get("/races/does-not-exist/readings.xlsx").status_code == 404
    assert client.get("/races/does-not-exist/results.xlsx").status_code == 404

    second = client.post("/races", json={"name": "Lauf 2", "total_laps": 2}).json()
    client.post(f"/races/{second['id']}/activate")
    race_id = [r["id"] for r in client.get("/races").json()["items"] if r["id"] != second["id"]][0]
    assert client.get(f"/races/{race_id}/results.xlsx").status_code == 409


def test_readings_export_scales(app_client):
    """A full event day must not take minutes or hundreds of megabytes."""
    client, app_module = app_client
    _seed(client, app_module)
    race_id = client.get("/race").json()["id"]

    rows = [
        (race_id, f"T{i % 200:04d}", "arrive",
         _iso(START + timedelta(seconds=100 + i)), (i % 2) + 1, -50.0, "00179E01")
        for i in range(20000)
    ]
    app_module.storage._conn.executemany(
        "INSERT OR IGNORE INTO tag_events "
        "(race_id, tag_id, event_type, timestamp, antenna, rssi, reader_serial) "
        "VALUES (?, ?, ?, ?, ?, ?, ?);",
        rows,
    )
    app_module.storage._conn.commit()

    res = client.get(f"/races/{race_id}/readings.xlsx")
    assert res.status_code == 200
    _, _, _, sheet_rows = _sheet_rows(res.content, "Lesungen")
    assert len(sheet_rows) == 20005


def test_filename_sanitising():
    from exports_xlsx import sanitize_filename

    assert sanitize_filename("Racetag-Ergebnis-Straßenrennen Süd-2026-09-24") == (
        "Racetag-Ergebnis-Strassenrennen-Sued-2026-09-24.xlsx"
    )
    assert sanitize_filename("///").endswith("export.xlsx")


# ---------------------------------------------------------------------------
# Tag inventory, start list, coupling delete (field report 2026-09-26)
# ---------------------------------------------------------------------------

def test_tags_workbook_keeps_leading_zeros_as_text(app_client):
    """Excel turns a digits-only EPC into a number and the re-import then
    couples a tag that does not exist. The cell must stay text."""
    client, app_module = app_client
    client.post("/riders", json={"tag_id": "000000009969", "bib": "7", "name": "Mit Nullen"})
    _post_arrive(client, "000000009969", START + timedelta(seconds=30))

    wb = load_workbook(io.BytesIO(client.get("/tags.xlsx").content))
    ws = wb["Tags"]
    assert ws["A1"].value == "tag_id"
    cell = ws["A2"]
    assert cell.value == "000000009969"
    assert isinstance(cell.value, str)
    assert cell.number_format == "@"
    assert ws["B2"].value == "7"
    assert ws["C2"].value == "Mit Nullen"


def test_tags_workbook_finds_the_name_in_another_race(app_client):
    """Couplings are per race. A tag waved in a fresh race still shows the
    number it was coupled to before, plus the race it came from."""
    client, app_module = app_client
    client.post("/riders", json={"tag_id": "AAA1", "bib": "7", "name": "Jürgen Öhler"})

    second = client.post("/races", json={"name": "Lauf 2", "total_laps": 2}).json()
    client.post(f"/races/{second['id']}/activate")
    _post_arrive(client, "AAA1", START + timedelta(hours=1))

    _, _, header, rows = _sheet_rows(client.get("/tags.xlsx").content, "Tags")
    cols = {name: i for i, name in enumerate(header)}
    row = next(r for r in rows if r[cols["tag_id"]] == "AAA1")
    assert row[cols["bib"]] == "7"
    assert row[cols["name"]] == "Jürgen Öhler"
    assert row[cols["Name aus Rennen"]] == "Default race"


def test_tags_workbook_prefers_the_coupling_of_this_race(app_client):
    client, app_module = app_client
    client.post("/riders", json={"tag_id": "AAA1", "bib": "7", "name": "Erstes Rennen"})

    second = client.post("/races", json={"name": "Lauf 2", "total_laps": 2}).json()
    client.post(f"/races/{second['id']}/activate")
    client.post("/riders", json={"tag_id": "AAA1", "bib": "99", "name": "Zweites Rennen"})
    _post_arrive(client, "AAA1", START + timedelta(hours=1))

    _, _, header, rows = _sheet_rows(client.get("/tags.xlsx").content, "Tags")
    cols = {name: i for i, name in enumerate(header)}
    row = next(r for r in rows if r[cols["tag_id"]] == "AAA1")
    assert row[cols["bib"]] == "99"
    assert row[cols["Name aus Rennen"]] is None  # no foreign race involved


def test_startlist_workbook_lists_riders_without_readings(app_client):
    client, app_module = app_client
    _seed(client, app_module)
    client.post("/riders", json={"tag_id": "DDD4", "bib": "99", "name": "Ohne Lesung"})

    res = client.get("/startlist.xlsx")
    assert res.status_code == 200
    wb, ws, header, rows = _sheet_rows(res.content, "Startliste")
    assert wb.sheetnames == ["Startliste", "Info"]
    cols = {name: i for i, name in enumerate(header)}
    by_bib = {r[cols["Startnummer"]]: r for r in rows}
    assert by_bib["99"][cols["Lesungen in diesem Rennen"]] == 0
    assert by_bib["7"][cols["Lesungen in diesem Rennen"]] == 3
    assert by_bib["7"][cols["tag_id"]] == "AAA1"
    assert ws.cell(row=2, column=cols["tag_id"] + 1).number_format == "@"


def test_startlist_workbook_works_for_an_inactive_race(app_client):
    client, app_module = app_client
    _seed(client, app_module)
    race_id = client.get("/race").json()["id"]
    second = client.post("/races", json={"name": "Lauf 2", "total_laps": 2}).json()
    client.post(f"/races/{second['id']}/activate")

    res = client.get(f"/races/{race_id}/startlist.xlsx")
    assert res.status_code == 200
    _, _, header, rows = _sheet_rows(res.content, "Startliste")
    cols = {name: i for i, name in enumerate(header)}
    assert {r[cols["Startnummer"]] for r in rows} == {"7", "12"}
    assert client.get("/races/does-not-exist/startlist.xlsx").status_code == 404
    assert client.get("/races/does-not-exist/tags.xlsx").status_code == 404


def test_deleting_a_coupling_removes_the_standings_row(app_client):
    """"Kopplung löschen" for a tag coupled to a wrong number: the rider has
    to disappear from the standings right away, while the readings stay."""
    client, app_module = app_client
    _seed(client, app_module)
    before = client.get("/classification").json()["standings"]
    assert any(p["tag_id"] == "AAA1" for p in before)

    assert client.delete("/riders/AAA1").status_code == 204

    after = client.get("/classification").json()["standings"]
    assert not any(p["tag_id"] == "AAA1" for p in after)
    assert client.get("/riders/AAA1").status_code == 404
    # The readings survive for the audit trail.
    _, _, header, rows = _sheet_rows(client.get("/readings.xlsx").content, "Lesungen")
    cols = {name: i for i, name in enumerate(header)}
    assert sum(1 for r in rows if r[cols["tag_id".replace("tag_id", "Tag-ID")]] == "AAA1") == 3
    assert client.delete("/riders/AAA1").status_code == 404


def test_deleting_a_race_removes_its_riders_and_readings(app_client):
    """Erik created a race by mistake; deleting it must not leave orphans."""
    client, app_module = app_client
    _seed(client, app_module)
    doomed = client.post("/races", json={"name": "Versehen", "total_laps": 2}).json()
    client.post(f"/races/{doomed['id']}/activate")
    client.post("/riders", json={"tag_id": "ZZZ9", "bib": "5", "name": "Falsch"})
    _post_arrive(client, "ZZZ9", START + timedelta(hours=2))

    # The active race is protected.
    assert client.delete(f"/races/{doomed['id']}").status_code == 409

    first = [r["id"] for r in client.get("/races").json()["items"] if r["id"] != doomed["id"]][0]
    client.post(f"/races/{first}/activate")
    assert client.delete(f"/races/{doomed['id']}").status_code == 204

    assert client.get(f"/races/{doomed['id']}").status_code == 404
    rows = app_module.storage._conn.execute(
        "SELECT (SELECT COUNT(*) FROM riders WHERE race_id = ?) AS riders, "
        "       (SELECT COUNT(*) FROM tag_events WHERE race_id = ?) AS events;",
        (doomed["id"], doomed["id"]),
    ).fetchone()
    assert rows["riders"] == 0
    assert rows["events"] == 0


def test_renaming_a_race_keeps_its_data(app_client):
    client, app_module = app_client
    _seed(client, app_module)
    race_id = client.get("/race").json()["id"]

    res = client.patch(f"/races/{race_id}", json={"name": "10 km Lauf"})
    assert res.status_code == 200
    assert res.json()["name"] == "10 km Lauf"
    assert client.get("/race").json()["name"] == "10 km Lauf"
    assert len(client.get("/classification").json()["standings"]) == 2
